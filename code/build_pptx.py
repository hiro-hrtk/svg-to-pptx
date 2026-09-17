"""出力層。全体の変換パイプライン本体。SVG(図形+テキスト) -> PPTXネイティブオブジェクト。

使い方:
    uv run python code/build_pptx.py --session <name> \
        --no-text <no_text_or_single.svg> --text-only <text_only_or_single.svg>

単一SVG(図形+テキスト混在)の場合は --no-text と --text-only に同じパスを渡す。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree as LET

import concat_text
import fit_text
import mapping as mapping_mod
import svg_elements

CODE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CODE_DIR.parent

EMU_PER_PX_BASE = 1.0  # placeholder, real scale computed per-document


# ── 座標変換 ────────────────────────────────────────────────

class CoordMapper:
    def __init__(self, svg_w, svg_h, slide_w_emu, slide_h_emu):
        self.svg_w = svg_w
        self.svg_h = svg_h
        self.scale = min(slide_w_emu / svg_w, slide_h_emu / svg_h)
        drawn_w = svg_w * self.scale
        drawn_h = svg_h * self.scale
        self.offset_x_emu = int((slide_w_emu - drawn_w) / 2)
        self.offset_y_emu = int((slide_h_emu - drawn_h) / 2)

    def point_emu(self, x_px, y_px):
        return (int(self.offset_x_emu + x_px * self.scale),
                int(self.offset_y_emu + y_px * self.scale))

    def len_emu(self, v_px):
        return int(v_px * self.scale)

    def px_to_pt(self, v_px):
        return v_px * self.scale / 12700.0


def _read_svg_size(svg_path):
    tree = ET.parse(svg_path)
    root = tree.getroot()
    vb = root.get("viewBox")
    if vb:
        parts = [float(p) for p in vb.replace(",", " ").split()]
        return parts[2], parts[3]
    return float(root.get("width", 1600)), float(root.get("height", 900))


# ── 色・スタイルのXMLヘルパー ────────────────────────────────

def _sub(parent, tag, **attrs):
    el = LET.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _apply_alpha_to_srgb(spPr_child_finder, alpha):
    """spPr配下の a:solidFill/a:srgbClr に a:alpha を付与する。"""
    if alpha >= 0.999:
        return
    solid_fill = spPr_child_finder
    if solid_fill is None:
        return
    srgb = solid_fill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    _sub(srgb, "a:alpha", val=str(int(max(0.0, min(1.0, alpha)) * 100000)))


def _set_dash(shape):
    ln = shape.line._get_or_add_ln()
    _sub(ln, "a:prstDash", val="dash")


def apply_shape_style(shape, s: svg_elements.ShapeElem, cm: CoordMapper, paint_fill=True, has_fill_attr=True):
    """paint_fill: このシェイプに実際に色を塗るか(閉じていないpath/polylineはFalse)。
    has_fill_attr: shapeオブジェクトが`.fill`属性を持つか(Connectorは持たない)。
    has_fill_attr=Trueの場合、paint_fillがFalseでもfill.background()を明示的に呼び、
    PowerPointのデフォルトテーマ塗り(青)が残らないようにする。
    """
    spPr = shape._element.spPr

    if has_fill_attr:
        if paint_fill and s.fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(s.fill)
            _apply_alpha_to_srgb(spPr.find(qn("a:solidFill")), s.opacity)
        else:
            shape.fill.background()

    if s.stroke:
        width_px = s.stroke_width if s.stroke_width > 0 else 1.0
        shape.line.color.rgb = RGBColor.from_string(s.stroke)
        shape.line.width = Pt(max(cm.px_to_pt(width_px), 0.25))
        ln = shape._element.spPr.find(qn("a:ln"))
        _apply_alpha_to_srgb(ln.find(qn("a:solidFill")) if ln is not None else None, s.opacity)
        if s.dasharray:
            _set_dash(shape)
    else:
        shape.line.fill.background()


# ── 図形の描画 ──────────────────────────────────────────────

def draw_rect(slide, s: svg_elements.ShapeElem, cm: CoordMapper):
    left, top = cm.point_emu(s.x, s.y)
    w, h = cm.len_emu(s.w), cm.len_emu(s.h)
    shape_type = MSO_SHAPE.RECTANGLE
    if s.rx > 0:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, w, h)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE and s.w > 0 and s.h > 0:
        adj = min(0.5, s.rx / min(s.w, s.h))
        try:
            shape.adjustments[0] = adj
        except (IndexError, ValueError):
            pass
    apply_shape_style(shape, s, cm)
    return shape


def draw_oval(slide, s: svg_elements.ShapeElem, cm: CoordMapper):
    left, top = cm.point_emu(s.x, s.y)
    w, h = cm.len_emu(s.w), cm.len_emu(s.h)
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, w, h)
    apply_shape_style(shape, s, cm)
    return shape


def draw_line(slide, s: svg_elements.ShapeElem, cm: CoordMapper):
    (x1, y1), (x2, y2) = s.points[0], s.points[1]
    ex1, ey1 = cm.point_emu(x1, y1)
    ex2, ey2 = cm.point_emu(x2, y2)
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ex1, ey1, ex2, ey2)
    apply_shape_style(shape, s, cm, has_fill_attr=False)
    return shape


def draw_freeform(slide, s: svg_elements.ShapeElem, cm: CoordMapper):
    pts = s.points
    if len(pts) < 2:
        return None
    fb = slide.shapes.build_freeform(start_x=pts[0][0], start_y=pts[0][1], scale=cm.scale)
    fb.add_line_segments(pts[1:], close=s.closed)
    shape = fb.convert_to_shape(origin_x=cm.offset_x_emu, origin_y=cm.offset_y_emu)
    apply_shape_style(shape, s, cm, paint_fill=s.closed)
    return shape


def draw_image(slide, s: svg_elements.ShapeElem, cm: CoordMapper):
    if not s.image_data:
        return None
    import io
    left, top = cm.point_emu(s.x, s.y)
    w, h = cm.len_emu(s.w), cm.len_emu(s.h)
    return slide.shapes.add_picture(io.BytesIO(s.image_data), left, top, w, h)


def draw_shape(slide, s: svg_elements.ShapeElem, cm: CoordMapper):
    if s.kind == "rect":
        return draw_rect(slide, s, cm)
    if s.kind in ("circle", "ellipse"):
        return draw_oval(slide, s, cm)
    if s.kind == "line":
        return draw_line(slide, s, cm)
    if s.kind in ("polyline", "polygon", "path"):
        return draw_freeform(slide, s, cm)
    if s.kind == "image":
        return draw_image(slide, s, cm)
    return None


# ── テキストの配置 ──────────────────────────────────────────

def _is_bold(weight: str) -> bool:
    weight = (weight or "").strip().lower()
    if weight == "bold":
        return True
    if weight.isdigit():
        return int(weight) >= 600
    return False


ANCHOR_ALIGN = {
    "start": PP_ALIGN.LEFT,
    "middle": PP_ALIGN.CENTER,
    "end": PP_ALIGN.RIGHT,
}


def _set_run_lang_and_charset(run):
    """PowerPointのスペルチェック誤検出(日本語への赤波線)を抑制するためlang/altLangを、
    Meiryo UIのUnicode範囲を明示するためa:latinのcharsetを設定する。
    """
    rPr = run._r.get_or_add_rPr()
    run.font.language_id = MSO_LANGUAGE_ID.JAPANESE
    rPr.set("altLang", "ja-JP")
    latin = rPr.get_or_add_latin()
    latin.set("charset", "0")


def place_text_block(slide, block: concat_text.TextBlock, shapes_by_id, cm: CoordMapper, config):
    inner_margin_px = config.get("text_box", {}).get("inner_margin_px", 4)
    line_spacing = config.get("text_box", {}).get("line_spacing", 1.15)
    floating_bg = config.get("text_box", {}).get("floating_bg_fill", "FFFFFF")

    container = shapes_by_id.get(block.shape_id) if block.shape_id else None
    if block.contained and container is not None:
        cl, ct, cw, ch = container.bbox
        content_width_px = max(cw - 2 * inner_margin_px, 10)
        content_height_px = max(ch - 2 * inner_margin_px, 10)
    else:
        content_width_px = None
        content_height_px = None

    fit = fit_text.fit_block(block, content_width_px, content_height_px, config)

    small_container_area_px2 = config.get("text_box", {}).get("small_container_area_px2", 60000)
    is_small_container = (
        block.contained and container is not None
        and content_width_px is not None
        and container.area <= small_container_area_px2
        # anchor='middle'(元々中央揃え意図)のみ対象。start/endはコンテナ内の他要素(アイコン等)を
        # 避けるために意図的にその位置にあることが多く、強制的に幅いっぱいに広げると重なる。
        and block.anchor == "middle"
    )

    box_width_px = fit.box_width_px if fit.box_width_px > 0 else fit.font_size_px
    align_override = None
    if is_small_container:
        # 小さいコンテナ内はanchor_xに関わらずコンテナ幅いっぱいに広げて中央揃えにする(ADR-015)。
        # anchor_x基準だとコンテナ幅より不必要に狭い/ずれた位置に出ることがあったため、
        # 小さいコンテナに限りコンテナ自身の幅を信頼する。
        box_l_px = container.bbox[0] + inner_margin_px
        box_width_px = content_width_px
        align_override = PP_ALIGN.CENTER
    elif block.anchor == "start":
        box_l_px = block.left
    elif block.anchor == "end":
        box_l_px = block.left - box_width_px
    else:
        box_l_px = block.left - box_width_px / 2
    if (
        block.contained and container is not None
        and len(block.paragraphs) == 1 and len(fit.lines) > 1
    ):
        # 元は1行のテキストだったが、収まりきらずfit_text側で複数行に折り返された場合。
        # block.topは1行想定の位置のままなので、コンテナの縦範囲からはみ出してしまう。
        # (concat済みの元々複数行のブロックはblock.topをそのまま使うためここでは対象外)
        ct_box_top = container.bbox[1] + inner_margin_px
        ct_box_h = max(container.bbox[3] - 2 * inner_margin_px, fit.box_height_px)
        box_t_px = ct_box_top + max(0.0, (ct_box_h - fit.box_height_px) / 2)
    else:
        box_t_px = block.top

    left, top = cm.point_emu(box_l_px, box_t_px)
    width = max(cm.len_emu(box_width_px), Emu(1))
    height = max(cm.len_emu(fit.box_height_px), Emu(1))

    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    # 折り返しはfit_text側で行済み(実測px幅ベース)。PPTX側の再折り返しに任せると
    # 疑似ボールド等でフォントメトリクスが変わった際に意図しない追加改行が起きるため無効化する。
    tf.word_wrap = False
    tf.auto_size = None
    small_margin = Emu(Pt(1))
    tf.margin_left = tf.margin_right = small_margin
    tf.margin_top = tf.margin_bottom = small_margin

    # 白背景(線を隠す意匠の代替, ADR-004)は「近くに図形がある浮遊テキスト」にのみ適用する。
    # 近傍に何の図形も無い孤立テキスト(タイトル等)にまで敷くと不要な白い矩形が
    # 目立ってしまうため対象外にする。
    if not block.contained and block.shape_id is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(floating_bg)
    else:
        box.fill.background()
    box.line.fill.background()

    font_pt = max(cm.px_to_pt(fit.font_size_px), 1.0)
    align = align_override or ANCHOR_ALIGN.get(block.anchor, PP_ALIGN.LEFT)
    bold = _is_bold(block.font_weight)

    # PowerPointの「単一行間隔」はフォント内部メトリクスに依存し、font_sizeより
    # かなり大きくなることがある(Meiryoなど)。倍率指定ではなく厳密なPt値を指定して、
    # 隣接する独立ブロック同士が視覚的に重ならないようにする。
    line_height_pt = Pt(font_pt * (1.0 if len(fit.lines) <= 1 else line_spacing))
    for i, line_text in enumerate(fit.lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line_height_pt
        run = para.add_run()
        run.text = line_text
        run.font.size = Pt(font_pt)
        run.font.bold = bold
        run.font.name = config.get("font", {}).get("default_name", "Meiryo UI")
        run.font.color.rgb = RGBColor.from_string(block.fill if block.fill else "000000")
        _set_run_lang_and_charset(run)

    return box, fit


# ── パイプライン本体 ────────────────────────────────────────

def load_config():
    with open(CODE_DIR / "config.yaml", "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def build(no_text_svg: str, text_only_svg: str, session: str, output_dir: Path):
    config = load_config()

    svg_w, svg_h = _read_svg_size(no_text_svg)

    shapes = svg_elements.parse_shapes(no_text_svg)
    texts = svg_elements.parse_texts(text_only_svg)

    mapping, contained_ids = mapping_mod.build_mapping(shapes, texts, config, svg_size=(svg_w, svg_h))
    blocks = concat_text.build_text_blocks(texts, mapping, contained_ids, shapes, config)
    shapes_by_id = {s.id: s for s in shapes}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    cm = CoordMapper(svg_w, svg_h, prs.slide_width, prs.slide_height)

    for s in shapes:
        try:
            draw_shape(slide, s, cm)
        except Exception as e:  # noqa: BLE001
            print(f"[WARN] shape {s.id} ({s.kind}) の描画に失敗: {e}", file=sys.stderr)

    mapping_debug = []
    for block in blocks:
        try:
            _, fit = place_text_block(slide, block, shapes_by_id, cm, config)
        except Exception as e:  # noqa: BLE001
            print(f"[WARN] text block {block.id} の配置に失敗: {e}", file=sys.stderr)
            continue
        mapping_debug.append({
            "block_id": block.id,
            "shape_id": block.shape_id,
            "contained": block.contained,
            "anchor": block.anchor,
            "left_px": block.left,
            "top_px": block.top,
            "font_size_px_final": fit.font_size_px,
            "paragraphs": block.paragraphs,
        })

    output_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = output_dir / f"{session}_native.pptx"
    mapping_path = output_dir / f"{session}_mapping.json"

    prs.save(str(pptx_path))
    with open(mapping_path, "w", encoding="utf-8") as f:
        json.dump(mapping_debug, f, ensure_ascii=False, indent=2)

    return pptx_path, mapping_path


def main():
    parser = argparse.ArgumentParser(description="SVG -> PPTX native converter")
    parser.add_argument("--session", required=True)
    parser.add_argument("--no-text", required=True, dest="no_text")
    parser.add_argument("--text-only", required=True, dest="text_only")
    parser.add_argument("--output-dir", default=None)
    args = parser.parse_args()

    output_dir = Path(args.output_dir) if args.output_dir else PROJECT_ROOT / "output" / args.session

    pptx_path, mapping_path = build(args.no_text, args.text_only, args.session, output_dir)
    print(f"PPTX: {pptx_path}")
    print(f"Mapping: {mapping_path}")


if __name__ == "__main__":
    main()
